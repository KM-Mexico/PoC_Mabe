from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

from app.models import DocumentStatus


class ExtractionError(ValueError):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


@dataclass(frozen=True)
class ExtractedChunk:
    content: str
    page_number: int | None = None
    sheet_name: str | None = None
    cell_range: str | None = None


@dataclass(frozen=True)
class ExtractionResult:
    media_type: str
    status: DocumentStatus
    chunks: list[ExtractedChunk]


def split_text(content: str, limit: int = 2_000) -> list[str]:
    normalized = "\n".join(line.rstrip() for line in content.replace("\r\n", "\n").replace("\r", "\n").split("\n")).strip()
    if not normalized:
        return []
    pieces: list[str] = []
    remaining = normalized
    while len(remaining) > limit:
        boundary = remaining.rfind("\n", 0, limit + 1)
        if boundary < limit // 2:
            boundary = remaining.rfind(" ", 0, limit + 1)
        if boundary < limit // 2:
            boundary = limit
        pieces.append(remaining[:boundary].strip())
        remaining = remaining[boundary:].strip()
    if remaining:
        pieces.append(remaining)
    return pieces


class DocumentExtractor:
    def extract(self, original_name: str, content: bytes) -> ExtractionResult:
        extension = Path(original_name).suffix.lower()
        if extension in {".txt", ".md"}:
            return self._extract_text(extension, content)
        if extension == ".pdf":
            return self._extract_pdf(content)
        if extension == ".xlsx":
            return self._extract_excel(content)
        raise ExtractionError("unsupported_type", "Only TXT, MD, PDF and XLSX files are supported.")

    def _extract_text(self, extension: str, content: bytes) -> ExtractionResult:
        if b"\x00" in content:
            raise ExtractionError("invalid_signature", "The file does not contain valid plain text.")
        try:
            text = content.decode("utf-8")
        except UnicodeDecodeError as error:
            raise ExtractionError("invalid_utf8", "Text files must use UTF-8 encoding.") from error
        chunks = [ExtractedChunk(piece) for piece in split_text(text)]
        if not chunks:
            raise ExtractionError("empty_document", "The text file does not contain useful content.")
        media_type = "text/markdown" if extension == ".md" else "text/plain"
        return ExtractionResult(media_type, DocumentStatus.READY, chunks)

    def _extract_pdf(self, content: bytes) -> ExtractionResult:
        if not content.startswith(b"%PDF-"):
            raise ExtractionError("invalid_signature", "The file signature does not match PDF.")
        try:
            reader = PdfReader(BytesIO(content), strict=True)
            if reader.is_encrypted:
                raise ExtractionError("protected_document", "Password-protected PDFs are not supported.")
            chunks = [
                ExtractedChunk(piece, page_number=page_number)
                for page_number, page in enumerate(reader.pages, start=1)
                for piece in split_text(page.extract_text() or "")
            ]
        except ExtractionError:
            raise
        except Exception as error:
            raise ExtractionError("corrupt_pdf", "The PDF is corrupt or cannot be read.") from error
        status = DocumentStatus.READY if chunks else DocumentStatus.REQUIRES_OCR
        return ExtractionResult("application/pdf", status, chunks)

    def _extract_excel(self, content: bytes) -> ExtractionResult:
        if not content.startswith(b"PK"):
            raise ExtractionError("invalid_signature", "The file signature does not match XLSX.")
        try:
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
            chunks: list[ExtractedChunk] = []
            for worksheet in workbook.worksheets:
                for row in worksheet.iter_rows():
                    populated = [cell for cell in row if cell.value not in (None, "")]
                    if not populated:
                        continue
                    rendered = " | ".join(f"{cell.coordinate}: {cell.value}" for cell in populated)
                    cell_range = f"{populated[0].coordinate}:{populated[-1].coordinate}"
                    chunks.append(ExtractedChunk(rendered, sheet_name=worksheet.title, cell_range=cell_range))
            workbook.close()
        except Exception as error:
            raise ExtractionError("corrupt_xlsx", "The XLSX file is corrupt or cannot be read.") from error
        if not chunks:
            raise ExtractionError("empty_document", "The workbook does not contain useful values.")
        return ExtractionResult("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", DocumentStatus.READY, chunks)
